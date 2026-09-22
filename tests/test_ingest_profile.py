from __future__ import annotations

import importlib.util
import json
from pathlib import Path

import pymupdf
from docx import Document


def load_module(root: Path):
    path = root / "scripts" / "ingest_profile.py"
    spec = importlib.util.spec_from_file_location("ingest_profile", path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    import sys
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def test_extract_docx_and_pdf(tmp_path):
    repo_root = Path(__file__).resolve().parents[1]
    mod = load_module(repo_root)

    source_dir = tmp_path / "sources"
    output_dir = tmp_path / "normalized"
    source_dir.mkdir()

    docx_path = source_dir / "resume.docx"
    doc = Document()
    doc.add_paragraph("Data analyst with Python experience")
    doc.add_table(rows=1, cols=2)
    doc.tables[0].cell(0, 0).text = "Degree"
    doc.tables[0].cell(0, 1).text = "MSc"
    doc.save(docx_path)

    pdf_path = source_dir / "resume.pdf"
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Machine learning researcher")
    pdf.save(pdf_path)
    pdf.close()

    index = tmp_path / "SOURCE_INDEX.json"
    evidence = tmp_path / "PROFILE_EVIDENCE.md"
    rc = mod.main([
        "--source-dir", str(source_dir),
        "--output-dir", str(output_dir),
        "--index", str(index),
        "--evidence", str(evidence),
    ])
    assert rc == 0

    payload = json.loads(index.read_text(encoding="utf-8"))
    assert payload["source_count"] == 2
    assert {x["source_type"] for x in payload["sources"]} == {".docx", ".pdf"}

    combined = "\n".join(p.read_text(encoding="utf-8") for p in output_dir.glob("*.md"))
    assert "Data analyst with Python experience" in combined
    assert "Machine learning researcher" in combined
    assert "SHA-256" in combined


def test_text_and_markdown_sources(tmp_path):
    repo_root = Path(__file__).resolve().parents[1]
    mod = load_module(repo_root)

    source_dir = tmp_path / "sources"
    output_dir = tmp_path / "normalized"
    source_dir.mkdir()
    (source_dir / "notes.txt").write_text("SQL and reporting", encoding="utf-8")
    (source_dir / "portfolio.md").write_text("# Portfolio\nForecasting project", encoding="utf-8")

    rc = mod.main([
        "--source-dir", str(source_dir),
        "--output-dir", str(output_dir),
        "--index", str(tmp_path / "index.json"),
        "--evidence", str(tmp_path / "evidence.md"),
    ])
    assert rc == 0
    combined = "\n".join(p.read_text(encoding="utf-8") for p in output_dir.glob("*.md"))
    assert "SQL and reporting" in combined
    assert "Forecasting project" in combined
