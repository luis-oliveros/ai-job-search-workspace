#!/usr/bin/env python3
"""Normalize career profile documents into auditable Markdown evidence.

Supported input formats: PDF, DOCX, TXT, MD.
PDF OCR is optional and requires pytesseract, Pillow, and the Tesseract binary.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

try:
    import pymupdf as fitz
except ImportError:  # PyMuPDF historically exposes the fitz module
    import fitz  # type: ignore

from docx import Document

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


@dataclass
class SourceResult:
    source_file: str
    source_type: str
    sha256: str
    extracted_at_utc: str
    extraction_method: str
    normalized_file: str
    characters: int
    warnings: list[str]


def utc_now() -> str:
    return datetime.now(timezone.utc).replace(microsecond=0).isoformat()


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def slugify(name: str) -> str:
    value = name.strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    return value.strip("_") or "source"


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.splitlines()]
    return "\n".join(lines).strip()


def ocr_page(page, lang: str) -> str:
    try:
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise RuntimeError(
            "OCR dependencies are not installed. Run: pip install -r requirements-ocr.txt"
        ) from exc

    matrix = fitz.Matrix(2.0, 2.0)
    pix = page.get_pixmap(matrix=matrix, alpha=False)
    image = Image.open(io.BytesIO(pix.tobytes("png")))
    return pytesseract.image_to_string(image, lang=lang)


def extract_pdf(path: Path, ocr_mode: str, ocr_lang: str, min_chars: int) -> tuple[str, str, list[str]]:
    warnings: list[str] = []
    sections: list[str] = []
    used_ocr = False
    doc = fitz.open(path)

    try:
        for index, page in enumerate(doc, start=1):
            direct = clean_text(page.get_text("text") or "")
            page_text = direct
            method = "text-layer"

            needs_ocr = ocr_mode == "force" or (
                ocr_mode == "auto" and len(re.sub(r"\s+", "", direct)) < min_chars
            )

            if needs_ocr:
                try:
                    ocr_text = clean_text(ocr_page(page, ocr_lang))
                    if ocr_text:
                        page_text = ocr_text
                        method = "ocr"
                        used_ocr = True
                    else:
                        warnings.append(f"Page {index}: OCR returned no text.")
                except Exception as exc:
                    warnings.append(f"Page {index}: OCR failed: {exc}")

            if not page_text:
                warnings.append(f"Page {index}: no extractable text.")
                page_text = "[NO EXTRACTABLE TEXT]"

            sections.append(f"## Page {index}\n\nExtraction: `{method}`\n\n{page_text}")
    finally:
        doc.close()

    method = "pdf-text+ocr" if used_ocr else "pdf-text"
    return "\n\n".join(sections), method, warnings


def extract_docx(path: Path) -> tuple[str, str, list[str]]:
    doc = Document(path)
    sections: list[str] = []
    warnings: list[str] = []

    paragraphs = []
    for i, paragraph in enumerate(doc.paragraphs, start=1):
        text = clean_text(paragraph.text)
        if text:
            paragraphs.append(f"[{i}] {text}")
    if paragraphs:
        sections.append("## Paragraphs\n\n" + "\n\n".join(paragraphs))

    for t_index, table in enumerate(doc.tables, start=1):
        rows = []
        for r_index, row in enumerate(table.rows, start=1):
            cells = [clean_text(cell.text).replace("\n", " / ") for cell in row.cells]
            rows.append(f"Row {r_index}: " + " | ".join(cells))
        sections.append(f"## Table {t_index}\n\n" + "\n".join(rows))

    if not sections:
        warnings.append("DOCX contains no extractable paragraph or table text.")
        sections.append("[NO EXTRACTABLE TEXT]")

    return "\n\n".join(sections), "docx-xml", warnings


def extract_text_file(path: Path) -> tuple[str, str, list[str]]:
    warnings: list[str] = []
    try:
        text = path.read_text(encoding="utf-8")
        method = "utf-8"
    except UnicodeDecodeError:
        text = path.read_text(encoding="utf-8", errors="replace")
        method = "utf-8-replacement"
        warnings.append("Invalid UTF-8 bytes were replaced during extraction.")
    return clean_text(text), method, warnings


def extract_source(path: Path, ocr_mode: str, ocr_lang: str, min_chars: int):
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path, ocr_mode, ocr_lang, min_chars)
    if ext == ".docx":
        return extract_docx(path)
    if ext in {".txt", ".md"}:
        return extract_text_file(path)
    raise ValueError(f"Unsupported extension: {ext}")


def unique_output_path(output_dir: Path, source: Path, used: set[str]) -> Path:
    base = f"{slugify(source.stem)}_{source.suffix.lower().lstrip('.')}"
    candidate = base
    n = 2
    while candidate in used:
        candidate = f"{base}_{n}"
        n += 1
    used.add(candidate)
    return output_dir / f"{candidate}.md"


def discover_sources(source_dir: Path) -> list[Path]:
    if not source_dir.exists():
        return []
    return sorted(
        [p for p in source_dir.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED],
        key=lambda p: p.name.lower(),
    )


def write_normalized(path: Path, source: Path, digest: str, method: str, extracted_at: str, body: str, warnings: list[str]):
    warning_text = "\n".join(f"- {w}" for w in warnings) if warnings else "- None"
    content = (
        f"# Normalized source: {source.name}\n\n"
        f"- Original file: `{source.name}`\n"
        f"- File type: `{source.suffix.lower()}`\n"
        f"- SHA-256: `{digest}`\n"
        f"- Extracted at UTC: `{extracted_at}`\n"
        f"- Extraction method: `{method}`\n\n"
        f"## Extraction warnings\n\n{warning_text}\n\n"
        f"## Extracted content\n\n{body.strip()}\n"
    )
    path.write_text(content, encoding="utf-8")


def build_evidence_md(results: list[SourceResult]) -> str:
    lines = [
        "# Profile evidence index",
        "",
        "Generated by `scripts/ingest_profile.py`.",
        "",
        "This file indexes extracted evidence. It does not declare that every extracted statement is verified.",
        "",
    ]
    if not results:
        lines.append("No supported source documents were found.")
        return "\n".join(lines) + "\n"

    for item in results:
        lines += [
            f"## {item.source_file}",
            "",
            f"- Type: `{item.source_type}`",
            f"- SHA-256: `{item.sha256}`",
            f"- Extraction method: `{item.extraction_method}`",
            f"- Normalized file: `{item.normalized_file}`",
            f"- Extracted characters: `{item.characters}`",
            f"- Extracted at UTC: `{item.extracted_at_utc}`",
            "- Warnings:",
        ]
        if item.warnings:
            lines += [f"  - {w}" for w in item.warnings]
        else:
            lines.append("  - None")
        lines.append("")
    return "\n".join(lines)


def parse_args(argv: Iterable[str] | None = None) -> argparse.Namespace:
    root = Path(__file__).resolve().parents[1]
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=root / "profile_sources")
    parser.add_argument("--output-dir", type=Path, default=root / "profile" / "normalized_sources")
    parser.add_argument("--index", type=Path, default=root / "profile" / "SOURCE_INDEX.json")
    parser.add_argument("--evidence", type=Path, default=root / "profile" / "PROFILE_EVIDENCE.md")
    parser.add_argument("--ocr", choices=["off", "auto", "force"], default="off")
    parser.add_argument("--ocr-lang", default="eng")
    parser.add_argument("--min-pdf-text-chars", type=int, default=40)
    parser.add_argument("--clear", action="store_true", help="Remove existing normalized .md files before extraction.")
    return parser.parse_args(argv)


def main(argv: Iterable[str] | None = None) -> int:
    args = parse_args(argv)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.index.parent.mkdir(parents=True, exist_ok=True)
    args.evidence.parent.mkdir(parents=True, exist_ok=True)

    if args.clear:
        for p in args.output_dir.glob("*.md"):
            p.unlink()

    sources = discover_sources(args.source_dir)
    results: list[SourceResult] = []
    used_names: set[str] = set()

    for source in sources:
        digest = sha256_file(source)
        extracted_at = utc_now()
        body, method, warnings = extract_source(
            source,
            ocr_mode=args.ocr,
            ocr_lang=args.ocr_lang,
            min_chars=args.min_pdf_text_chars,
        )
        output_path = unique_output_path(args.output_dir, source, used_names)
        write_normalized(output_path, source, digest, method, extracted_at, body, warnings)

        try:
            normalized_display = output_path.relative_to(args.index.parent.parent).as_posix()
        except ValueError:
            normalized_display = output_path.as_posix()

        results.append(
            SourceResult(
                source_file=source.name,
                source_type=source.suffix.lower(),
                sha256=digest,
                extracted_at_utc=extracted_at,
                extraction_method=method,
                normalized_file=normalized_display,
                characters=len(body),
                warnings=warnings,
            )
        )

    index_payload = {
        "generated_at_utc": utc_now(),
        "source_count": len(results),
        "sources": [asdict(item) for item in results],
    }
    args.index.write_text(json.dumps(index_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    args.evidence.write_text(build_evidence_md(results), encoding="utf-8")

    print(f"Processed {len(results)} supported source document(s).")
    print(f"Index: {args.index}")
    print(f"Evidence: {args.evidence}")
    if not results:
        print(f"Add PDF, DOCX, TXT, or MD files to: {args.source_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
